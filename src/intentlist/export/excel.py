"""Batch and master spreadsheet delivery."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..db import audit
from ..models import (
    COUNTABLE_EMAIL_STATUSES,
    SIDECAR_EMAIL_STATUSES,
    Batch,
    BatchItem,
    Candidate,
    CandidateStatus,
    Capture,
    Category,
    EmailStatus,
    utcnow,
)

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")

COLUMNS = [
    ("#", 5), ("Full Name", 22), ("Email Address", 30), ("Email Verification Status", 14), ("Source URL", 45),
    ("Item Sought", 22), ("Exact Item Description", 60), ("Date Posted", 12), ("Category", 18),
    ("Evidence: name", 30), ("Evidence: email", 30), ("Confidence", 10), ("Review Flags", 24), ("Reviewed By", 12),
    ("Reviewed At", 18), ("Record ID", 10), ("Captured At", 18), ("Snapshot SHA-256", 20),
]

STATUS_DEFINITIONS = [
    ("valid", "Mailbox confirmed by the verification provider. Counts toward the batch."),
    ("catch_all", "Domain accepts all addresses, so the mailbox can't be confirmed. Delivered on a separate tab; not counted."),
    ("invalid", "Mailbox does not exist. Excluded."),
    ("risky", "Spam-trap, abuse or do-not-mail signals. Excluded."),
    ("disposable", "Temporary/disposable address. Excluded."),
    ("role_based", "Shared role address (info@, sales@). Excluded."),
    ("syntax_error", "Not a valid address format. Excluded."),
    ("domain_error", "Domain has no mail server. Excluded."),
    ("unknown", "Provider could not determine. Excluded."),
]


class NotEnoughRecords(RuntimeError):
    def __init__(self, available: int, required: int):
        super().__init__(f"only {available} approved, verified records available; {required} required")
        self.available, self.required = available, required


def _exported_keys(session: Session) -> set[tuple[str, str]]:
    rows = session.execute(select(Candidate.email_norm, Candidate.request_fingerprint)
                           .join(BatchItem, BatchItem.candidate_id == Candidate.id))
    return {(e, f) for e, f in rows}


def _ready(session: Session, category_id: int, statuses: set[EmailStatus]) -> list[Candidate]:
    return list(session.scalars(
        select(Candidate).where(Candidate.category_id == category_id, Candidate.status == CandidateStatus.approved,
                                Candidate.email_status.in_(statuses))
        .order_by(Candidate.reviewed_at, Candidate.id)))


def _dedupe_final(cands: list[Candidate], exported: set[tuple[str, str]]) -> list[Candidate]:
    """Belt and braces: one row per person+item, never anything already delivered in any batch."""
    out, seen = [], set(exported)
    for c in cands:
        key = (c.email_norm, c.request_fingerprint)
        if key in seen:
            continue
        seen.add(key)
        out.append(c)
    return out


def _style_sheet(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def _write_rows(ws: Worksheet, cands: list[Candidate], captures: dict[int, Capture], cat_names: dict[int, str]) -> None:
    ws.append([c[0] for c in COLUMNS])
    for n, c in enumerate(cands, start=1):
        cap = captures[c.capture_id]
        ev = c.evidence or {}
        ws.append([
            n, c.full_name, c.email, c.email_status.value if c.email_status else "", cap.url, c.item_sought,
            c.item_description, c.date_posted, cat_names.get(c.category_id, ""), ev.get("full_name", ""),
            ev.get("email", ""), round(c.confidence, 2), ", ".join(c.flags or []), c.reviewed_by,
            c.reviewed_at.replace(tzinfo=None) if c.reviewed_at else None, c.id,
            cap.fetched_at.replace(tzinfo=None) if cap.fetched_at else None, cap.content_sha256,
        ])
        row = ws.max_row
        url_cell = ws.cell(row=row, column=5)
        url_cell.hyperlink, url_cell.font = cap.url, LINK_FONT
        ws.cell(row=row, column=3).hyperlink = f"mailto:{c.email}"
        ws.cell(row=row, column=8).number_format = "yyyy-mm-dd"
        for col in (15, 17):
            ws.cell(row=row, column=col).number_format = "yyyy-mm-dd hh:mm"
        for col in (7, 10, 11):
            ws.cell(row=row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    _style_sheet(ws, [w for _, w in COLUMNS])


def _captures_for(session: Session, cands: list[Candidate]) -> dict[int, Capture]:
    ids = {c.capture_id for c in cands}
    return {c.id: c for c in session.scalars(select(Capture).where(Capture.id.in_(ids)))} if ids else {}


def export_batch(session: Session, category_slug: str, export_dir: Path, batch_size: int = 100,
                 allow_partial: bool = False, actor: str = "system") -> Batch:
    cat = session.scalar(select(Category).where(Category.slug == category_slug))
    if cat is None:
        raise ValueError(f"unknown category {category_slug!r}")
    exported = _exported_keys(session)
    ready = _dedupe_final(_ready(session, cat.id, COUNTABLE_EMAIL_STATUSES), exported)
    available_total = len(ready)
    if available_total < batch_size and not allow_partial:
        raise NotEnoughRecords(available_total, batch_size)
    if available_total == 0:
        raise NotEnoughRecords(0, batch_size)
    counted = ready[:batch_size]
    sidecar = _dedupe_final(_ready(session, cat.id, SIDECAR_EMAIL_STATUSES),
                            exported | {(c.email_norm, c.request_fingerprint) for c in counted})[:batch_size]

    number = (session.scalar(select(func.max(Batch.number)).where(Batch.category_id == cat.id)) or 0) + 1
    now = utcnow()
    out_dir = Path(export_dir) / cat.slug
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"batch_{number:03d}_{cat.slug}_{now:%Y%m%d}.xlsx"

    captures = _captures_for(session, counted + sidecar)
    names = {cat.id: cat.name}
    wb = Workbook()
    ws = wb.active
    ws.title = f"Batch {number}"
    _write_rows(ws, counted, captures, names)
    ws2 = wb.create_sheet("Catch-all (not counted)")
    _write_rows(ws2, sidecar, captures, names)

    summary = wb.create_sheet("Summary")
    rows = [
        ("Category", cat.name), ("Batch number", number), ("Generated (UTC)", now.replace(tzinfo=None)),
        ("Records counted (valid email)", len(counted)), ("Catch-all records (separate tab, not counted)", len(sidecar)),
        ("Approved valid records still available after this batch", available_total - len(counted)),
        ("Duplicate policy", "Checked against every record in every previously delivered batch, all categories."),
        ("Evidence policy", ("Every name, email and description appears verbatim in the captured source page; "
                             "the snapshot hash identifies the stored copy.")),
        ("Review policy", "Every record was approved by a named human reviewer."), ("", ""),
        ("Email status", "Meaning"), *STATUS_DEFINITIONS,
    ]
    for r in rows:
        summary.append(list(r))
    summary.column_dimensions["A"].width, summary.column_dimensions["B"].width = 48, 100
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    summary.cell(row=3, column=2).number_format = "yyyy-mm-dd hh:mm"
    wb.save(path)

    batch = Batch(category_id=cat.id, number=number, file_path=str(path), record_count=len(counted),
                  sidecar_count=len(sidecar), available_total=available_total)
    session.add(batch)
    session.flush()
    for c in counted:
        session.add(BatchItem(batch_id=batch.id, candidate_id=c.id, counted=True))
        c.status = CandidateStatus.exported
    for c in sidecar:
        session.add(BatchItem(batch_id=batch.id, candidate_id=c.id, counted=False))
        c.status = CandidateStatus.exported
    audit(session, "batch", batch.id, "exported", actor, file=str(path), counted=len(counted), sidecar=len(sidecar))
    session.flush()
    return batch


def export_master(session: Session, export_dir: Path) -> Path:
    """Every delivered record across all categories: the client's master de-duplication sheet."""
    rows = session.execute(
        select(Candidate, Batch.number, Category.name, BatchItem.counted)
        .join(BatchItem, BatchItem.candidate_id == Candidate.id).join(Batch, Batch.id == BatchItem.batch_id)
        .join(Category, Category.id == Batch.category_id).order_by(Category.name, Batch.number, Candidate.id)).all()
    cands = [r[0] for r in rows]
    captures = _captures_for(session, cands)
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    header = ["Category", "Batch", "Counted", "Full Name", "Email Address", "Email Verification Status", "Source URL",
              "Item Sought", "Exact Item Description", "Date Posted", "Record ID"]
    ws.append(header)
    for cand, number, cat_name, counted in rows:
        cap = captures[cand.capture_id]
        ws.append([cat_name, number, "yes" if counted else "no (catch-all)", cand.full_name, cand.email,
                   cand.email_status.value if cand.email_status else "", cap.url, cand.item_sought,
                   cand.item_description, cand.date_posted, cand.id])
        cell = ws.cell(row=ws.max_row, column=7)
        cell.hyperlink, cell.font = cap.url, LINK_FONT
        ws.cell(row=ws.max_row, column=10).number_format = "yyyy-mm-dd"
    _style_sheet(ws, [22, 7, 14, 22, 30, 14, 45, 22, 60, 12, 10])
    out = Path(export_dir) / f"master_{datetime.now():%Y%m%d_%H%M}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
