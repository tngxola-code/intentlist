"""Branch 10: batch delivery, master sheet, and the full cycle."""
import json

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from intentlist.export.excel import NotEnoughRecords, export_batch, export_master
from intentlist.models import Batch, BatchItem, Candidate, CandidateStatus, Capture, Category, EmailStatus
from intentlist.orchestration.pipeline import Pipeline, Services
from intentlist.review.service import apply_decision


def _approve_all(db, slug):
    cat = db.scalar(select(Category).where(Category.slug == slug))
    ids = db.scalars(select(Candidate.id).where(Candidate.category_id == cat.id,
                                                Candidate.status == CandidateStatus.pending_review)).all()
    for cid in ids:
        apply_decision(db, cid, "alice", "approve")
    db.commit()
    return len(ids)


def test_full_cycle_delivers_exactly_100_with_no_decoys(settings, corpus, db):
    svc = Services.from_settings(settings)
    Pipeline(db, svc).run("watches")
    _approve_all(db, "watches")
    batch = export_batch(db, "watches", settings.export_dir, 100)
    db.commit()
    assert batch.record_count == 100

    index = json.loads((settings.fixture_dir / "pages" / "index.json").read_text())
    rows = db.execute(select(Capture.url, Candidate.email_status, BatchItem.counted)
                      .join(Candidate, Candidate.capture_id == Capture.id)
                      .join(BatchItem, BatchItem.candidate_id == Candidate.id)).all()
    assert all(index[u]["kind"] in ("good", "repeat", "catchall") for u, _, _ in rows)
    assert all(st == EmailStatus.valid for _, st, counted in rows if counted)
    assert all(st == EmailStatus.catch_all for _, st, counted in rows if not counted)

    dupes = db.execute(select(Candidate.email_norm, Candidate.request_fingerprint, func.count())
                       .join(BatchItem, BatchItem.candidate_id == Candidate.id)
                       .group_by(Candidate.email_norm, Candidate.request_fingerprint).having(func.count() > 1)).all()
    assert dupes == []

    wb = load_workbook(batch.file_path)
    assert wb.sheetnames == ["Batch 1", "Catch-all (not counted)", "Summary"]
    ws = wb["Batch 1"]
    assert ws.max_row == 101 and ws.cell(2, 5).hyperlink is not None and ws.auto_filter.ref

    Pipeline(db, svc).run("watches")
    _approve_all(db, "watches")
    with pytest.raises(NotEnoughRecords):
        export_batch(db, "watches", settings.export_dir, 100)
    assert export_master(db, settings.export_dir).exists()


def test_delivered_request_is_not_redelivered_when_crossposted(settings, corpus, db):
    Pipeline(db, Services.from_settings(settings)).run("cars")
    _approve_all(db, "cars")
    export_batch(db, "cars", settings.export_dir, 100, allow_partial=True)
    db.commit()
    delivered = db.scalar(select(func.count()).select_from(BatchItem))
    cand = db.scalar(select(Candidate).join(BatchItem, BatchItem.candidate_id == Candidate.id).limit(1))
    cap = db.get(Capture, cand.capture_id)
    idx_path = settings.fixture_dir / "pages" / "index.json"
    index = json.loads(idx_path.read_text())
    index[cap.url + "-crosspost"] = dict(index[cap.url])
    idx_path.write_text(json.dumps(index))
    Pipeline(db, Services.from_settings(settings)).run("cars")
    _approve_all(db, "cars")
    export_batch(db, "cars", settings.export_dir, 100, allow_partial=True) if db.scalar(
        select(func.count()).select_from(Candidate).where(Candidate.status == CandidateStatus.approved,
                                                          Candidate.email_status == EmailStatus.valid)) else None
    rows = db.scalars(select(BatchItem).join(Candidate, Candidate.id == BatchItem.candidate_id)
                      .where(Candidate.email_norm == cand.email_norm,
                             Candidate.request_fingerprint == cand.request_fingerprint)).all()
    assert len(rows) == 1
    assert db.scalar(select(func.count()).select_from(BatchItem)) >= delivered


def test_export_refuses_short_batch_and_counts_only_valid(settings, corpus, db):
    Pipeline(db, Services.from_settings(settings)).run("cars")
    with pytest.raises(NotEnoughRecords):
        export_batch(db, "cars", settings.export_dir, 100)
    _approve_all(db, "cars")
    b = export_batch(db, "cars", settings.export_dir, 25)
    db.commit()
    assert b.record_count == 25
    counted = db.scalars(select(Candidate.email_status).join(BatchItem, BatchItem.candidate_id == Candidate.id)
                         .where(BatchItem.batch_id == b.id, BatchItem.counted.is_(True))).all()
    assert set(counted) == {EmailStatus.valid}
    assert export_batch(db, "cars", settings.export_dir, 25).number == 2
    assert db.scalar(select(func.count()).select_from(Batch)) == 2
